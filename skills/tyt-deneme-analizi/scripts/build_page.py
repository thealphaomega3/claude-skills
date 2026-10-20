#!/usr/bin/env python3
"""Extract a TYT tracking sheet (topic rows x exam columns, "X" = missed question)
into JSON and render the dashboard template.

Usage: build_page.py <xlsx> <out.html> [--style optik|defter] [--sheet "TYT AYRINTI"] [--json out.json]

Sheet layout this expects (one student's handwritten-style tracker):
- Several side-by-side blocks. Each block starts with a label column holding topic
  names; the columns to its right are exams. EVERY COLUMN IS ONE FULL TYT EXAM,
  shared by all subjects (a subject may be blank in a column = not recorded).
- Each subject is a run of topic rows followed by a "total" row whose exam cells hold
  the integer number of correct answers. An "X" (or "XX", "XXX") in a topic cell means
  that many questions from that topic were missed in that exam.
- Columns whose total cells hold floats are block averages and are skipped.
- Free text in a subject's cells (e.g. "3D", "APOT", "MSÜ") is the publisher.
"""
import sys, json, argparse, re
from collections import Counter

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl missing: python3 -m venv .v && .v/bin/pip install openpyxl, then run with .v/bin/python")

# Canonical topic names keyed by upper-cased, whitespace-collapsed sheet label.
CANON = {
    'SAYI KÜMELERİ(TEMEL KAVRAMLAR)': 'Sayı Kümeleri', 'SAYI ÇEŞİTLERİ(TEMEL KAVRAMLAR)': 'Sayı Çeşitleri',
    'FAKTÖRİYEL': 'Asal Sayılar / Faktöriyel', 'ASAL SAYILAR': 'Asal Sayılar / Faktöriyel',
    'RASYONEL ONDALIKLI SAYILAR': 'Rasyonel & Ondalık Sayılar',
    'KÜMELER,KARTEZYAN ÇARPIM': 'Kümeler', 'KÜMELER': 'Kümeler', 'RUTİN OLMAYAN PROBL': 'Rutin Olmayan Problemler',
    'ÖRÜNTÜ': 'Örüntü', 'MODÜLER ARTIMETIK, İŞLEM': 'Modüler Aritmetik', 'MODÜLER ARTİMETİK, İŞLEM': 'Modüler Aritmetik',
    'YAMUK': 'Dörtgenler', 'PARELEL KENAR-EŞKENAR DÖRTGEN-YAMUK': 'Dörtgenler', 'DELTAİD': 'Dörtgenler',
    'DİKDÖRTGEN': 'Dörtgenler', 'KARE': 'Dörtgenler', 'ÖZEL ÜÇGENLER': 'Özel Üçgenler',
    "NEWTON'UN HAREKET YASALARI": "Newton'un Hareket Yasaları", 'İŞ, ENERJİ VE GÜÇ': 'İş, Enerji ve Güç',
    'DAYANIKLILIK,ADEZYON,KOHEZYON': 'Dayanıklılık, Adezyon, Kohezyon', 'MIKNATIS,AKIM VE MANYETİK ALAN': 'Mıknatıs ve Manyetik Alan',
    'KİMYASALTÜRLERARASIETKİLEŞİMLER': 'Kimyasal Türler Arası Etkileşimler', 'KİMYANIN T.KANUNLARI,HESAPLAMAL.': 'Kimyanın Temel Kanunları',
    'ASİT,BAZ VE TUZLAR': 'Asit, Baz ve Tuzlar',
    'EKOLOJİ': 'Ekoloji / Ekosistem', 'EKOSİSTEM': 'Ekoloji / Ekosistem', 'CANLILAR DÜNYASI': 'Canlılar Dünyası',
    'CANLİLAR DÜNYASI': 'Canlılar Dünyası', 'CANLİLAR ALEMİ': 'Canlılar Dünyası', 'HÜCRE VE YAPISI': 'Hücre ve Yapısı',
    'CÜMLEDE ANLAM': 'Cümlede Anlam', 'CÜLENİN ÖĞ': 'Cümlenin Ögeleri', 'CÜMLENİNÖĞELERİ': 'Cümlenin Ögeleri',
    'TÜR': 'Metin Türleri', 'ANLATIM BİÇİMLERİ': 'Anlatım Biçimleri', 'CÜMLE TÜRLERİ': 'Cümle Türleri',
    'PARAGRAFTA YAPI': 'Paragrafta Yapı', 'FİİLİMSİ': 'Fiilimsi', 'FİİL': 'Fiil', 'TAMLAMALAR': 'Tamlamalar',
    'UYGARLIĞIN DOĞUŞU VE': 'Uygarlığın Doğuşu', 'BEYLİKTEN DEVL OSM.': 'Beylikten Devlete Osmanlı',
    'SELÇUKLU TÜRKİYESİ': 'Selçuklu Türkiyesi', 'ORTAÇAĞ DÜNYA': 'Orta Çağ Dünyası', 'DEVRİMLER ÇAI DEĞİ.': 'Değişim Çağı / Devrimler',
    '1.DÜNYA SAVAŞI VE MİLLİ MÜCADELE': 'I. Dünya Savaşı ve Milli Mücadele', 'TÜRK İNKLABI': 'Türk İnkılabı',
    'TOPRAK TİPLERİ': 'Toprak Tipleri', 'TÜRKİYR YERŞ': 'Türkiye Yer Şekilleri', 'SU,TOPRAK VE BİTKİ': 'Su, Toprak ve Bitki',
    'NÜFUS VE YERLEŞME': 'Nüfus ve Yerleşme', 'TÜRKİYEDE NÜFUS': "Türkiye'de Nüfus",
    'MÖ FELSEFE': 'MÖ Felsefe', 'VARLIK FELS': 'Varlık Felsefesi', 'M.S.2.YY-M.S.15.YY FELSEFESİ': 'MS 2.–15. yy Felsefesi',
    '15.YY-17.YY FELSEFESİ': '15.–17. yy Felsefesi', '18.YY-19.YY FELSEFESİ': '18.–19. yy Felsefesi', '20.YY FELSEFESİ': '20. yy Felsefesi',
    'HZ.MUHAMMED(S.A.V)': 'Hz. Muhammed (s.a.v.)',
    # labels seen in the newer "TAKİP TASLAK" sheet (word-problem types kept as their own topics there)
    'SAYI PROB': 'Sayı Problemleri', 'KESİR PROB': 'Kesir Problemleri', 'PERİYODİK PROB': 'Periyodik Problemler',
    'KAR ZARAR PROBL': 'Kâr-Zarar Problemleri', 'YÜZDE PEROBL': 'Yüzde Problemleri', 'KARIŞIM PROBLRMRLT': 'Karışım Problemleri',
    'YAŞ PROB': 'Yaş Problemleri', 'İŞÇİ PROB': 'İşçi Problemleri', 'HAREKET PROB': 'Hareket Problemleri',
    'POLİNOM': 'Polinomlar', 'POLİNOMLAR': 'Polinomlar', 'MİTOZ MAYOZ': 'Hücre Bölünmeleri', 'CANLILAR ALEM İ': 'Canlılar Dünyası',
}
# Word-problem sub-types get folded into one "Problemler" topic.
PROBLEM_SUBTYPES = {'PROBLEMLER', 'SAYI', 'KESİR', 'YAŞ', 'HIZ', 'İŞÇİ PROBLEMİ', 'ALIŞ VERİŞ', 'PERİYODİK', 'GRAFİK OKUMA', 'MANTIK MUHAKEME PROB'}
PUB = {'3D': '3D', '3 D': '3D', 'PRF': 'Paraf', 'PALME': 'Palme', 'PALM': 'Palme', 'TDR': 'TDR', 'ADAY': 'Aday',
       'APOT': 'Apotemi', 'APOTH': 'Apotemi', 'APTH': 'Apotemi', 'AOT': 'Apotemi', 'LİMİT': 'Limit', 'LMT': 'Limit',
       'ÖZ DE': 'Özdebir', 'ÖZDE': 'Özdebir', 'ÖZD': 'Özdebir', 'ÖZ D E': 'Özdebir', 'MSÜ': 'MSÜ', 'MEB': 'MEB', 'MEBİ': 'MEB',
       '345': '345', 'BSM': 'BSM', 'BLS': 'BLS', 'BRŞ': 'Branş', 'ÇAP': 'Çap', 'END': 'Endemik', 'AYD': 'Aydın', 'VIP': 'VIP',
       'ORJ': 'Orijinal', 'ACİL': 'Acil', 'KRALK': 'Kral Kızı', 'PAYL.': 'Paylaşım', 'BS': 'Bilgi Sarmal', 'RSY': 'RSY', 'ENS': 'ENS'}
# Subject detection: header label ("TYT FİZİK") or a telltale topic in the section.
SUBJECT_HINTS = [('Matematik', ['PROBLEMLER', 'OLASILIK']), ('Fizik', ['FİZİK', 'DALGALAR']), ('Kimya', ['KİMYA']),
                 ('Biyoloji', ['BİYOLOJİ', 'HÜCRE']), ('Türkçe', ['TÜRKÇE', 'PARAGRAF']), ('Tarih', ['TARİH']),
                 ('Coğrafya', ['COĞRAFYA', 'İKLİM']), ('Felsefe', ['FELSEFE']), ('Din', ['DİN', 'İBADET'])]
TYT_MAX = {'Türkçe': 40, 'Matematik': 40, 'Fizik': 7, 'Kimya': 7, 'Biyoloji': 6, 'Tarih': 5, 'Coğrafya': 5, 'Felsefe': 5, 'Din': 5}

def key(v):
    # Turkish-aware upper: Python maps i->I, but Turkish needs i->İ and ı->I
    s = str(v).replace('\xa0', ' ').replace('i', 'İ').replace('ı', 'I')
    return re.sub(r'\s+', ' ', s).strip().upper()
def isx(v): return isinstance(v, str) and v.strip() and set(v.strip().upper()) <= set('X')

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('out')
    ap.add_argument('--sheet', default=None); ap.add_argument('--json')
    ap.add_argument('--style', choices=['optik', 'defter'], default='optik')
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    name = a.sheet or next((n for n in wb.sheetnames if 'TYT' in n.upper() and 'AYR' in n.upper()), None)
    if not name: sys.exit(f'Sheet not found; pass --sheet. Sheets: {wb.sheetnames}')
    ws = wb[name]
    R, C = ws.max_row, ws.max_column
    cell = lambda r, c: ws.cell(r, c).value
    # label columns: many long text values
    labcols = [c for c in range(1, C + 1) if sum(1 for r in range(1, R + 1) if isinstance(cell(r, c), str) and len(cell(r, c).strip()) > 6 and not isx(cell(r, c))) > R * 0.4]
    blocks = [(lc, (labcols[i + 1] if i + 1 < len(labcols) else C + 1)) for i, lc in enumerate(labcols)]
    # exam columns actually in use (a draft sheet may fill only a few columns of one block)
    active = [c for lc, end in blocks for c in range(lc + 1, end) if any(cell(r, c) is not None for r in range(1, R + 1))]
    # a correct-answer count is 0..40; publisher codes like 345 are not
    isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and 0 <= v <= 40
    istext = lambda v: isinstance(v, str) and v.strip() != ''
    # total rows: counts in at least 3 active columns, more counts than text, no X marks
    # (a subject may cover only some exams, so no "most columns filled" rule)
    def is_total(r):
        n = sum(isnum(cell(r, c)) for c in active)
        return n >= min(3, len(active)) and n > sum(istext(cell(r, c)) for c in active) and not any(isx(cell(r, c)) for c in active)
    totals = [r for r in range(1, R + 1) if is_total(r)]
    avg = {c for lc, end in blocks for c in range(lc + 1, end) if any(isinstance(cell(r, c), float) and not float(cell(r, c)).is_integer() for r in totals)}
    subjects, prev = [], 0
    for tr in totals:
        rows = range(prev + 1, tr); prev = tr
        labels = ' | '.join(key(cell(r, lc)) for r in rows for lc, _ in blocks if cell(r, lc))
        sname = next((n for n, hints in SUBJECT_HINTS if any(h in labels for h in hints) and n not in [s['name'] for s in subjects]), None)
        if not sname: print(f'warn: could not name section ending row {tr}', file=sys.stderr); continue
        geo_from = next((r for r in rows if key(cell(r, 1) or '') == 'TYT GEOMETRİ'), None)
        tests, topics = [], []
        for bi, (lc, end) in enumerate(blocks):
            for c in range(lc + 1, end):
                t = cell(tr, c)
                if c in avg or not isinstance(t, int) or isinstance(t, bool): continue
                miss, pub = Counter(), None
                for r in rows:
                    v = cell(r, c)
                    if v is None: continue
                    if isx(v):
                        # a block may leave a label blank; borrow the same row's label from another block
                        lab = cell(r, lc) or next((cell(r, l2) for l2, _ in blocks if cell(r, l2)), None)
                        k = key(lab) if lab else ''
                        if k.startswith('TYT ') or k in ('TOPLAM', ''):
                            topic = 'Diğer'
                        else:
                            topic = 'Problemler' if (sname == 'Matematik' and k in PROBLEM_SUBTYPES) else CANON.get(k, str(lab).replace('\xa0', ' ').strip())
                        miss[topic] += len(v.strip())
                    elif pub is None:
                        pub = PUB.get(key(v))
                tests.append({'d': bi, 'x': c, 'c': t, 'p': pub, 'm': dict(miss)})
        for r in rows:
            lab = cell(r, 1)
            if not lab: continue
            k = key(lab)
            n = 'Problemler' if (sname == 'Matematik' and k in PROBLEM_SUBTYPES) else CANON.get(k, str(lab).strip())
            if n not in [t['n'] for t in topics]:
                topics.append({'n': n, 'g': 'Geometri' if (geo_from and r > geo_from) else None})
        mx = TYT_MAX.get(sname) or max(t['c'] for t in tests)
        subjects.append({'name': sname, 'max': mx, 'tests': tests, 'topics': topics})
    # summary for the agent to write findings from
    exams = {t['x'] for s in subjects for t in s['tests']}
    print(f'sheet={name!r} blocks={len(blocks)} exams(columns)={len(exams)} subjects={[s["name"] for s in subjects]}')
    for s in subjects:
        byd = {}
        for t in s['tests']: byd.setdefault(t['d'], []).append(t['c'])
        tm = Counter()
        for t in s['tests']: tm.update(t['m'])
        print(f"{s['name']:<10} n={len(s['tests']):>3} max={s['max']:>2} pct/dönem=" + ' '.join(f"D{d+1}:{sum(v)/len(v)/s['max']*100:.0f}" for d, v in sorted(byd.items())) + f"  top_miss={tm.most_common(4)}")
    data = json.dumps(subjects, ensure_ascii=False, separators=(',', ':'))
    if a.json: open(a.json, 'w').write(data)
    tpl = open(__file__.rsplit('/scripts/', 1)[0] + f'/assets/template-{a.style}.html', encoding='utf-8').read()
    open(a.out, 'w', encoding='utf-8').write(tpl.replace('__DATA__', data.replace('</', '<\\/')))
    print('wrote', a.out)

if __name__ == '__main__':
    main()
